import os
import tempfile
import docx
from app.models import TranslationResult
from app.output.docx import write_docx

def test_write_docx():
    results = [
        TranslationResult("chunk-1", "नमस्ते दुनिया", "hindi"),
        TranslationResult("chunk-2", "यह एक परीक्षण है", "hindi"),
    ]

    with tempfile.TemporaryDirectory() as tmpdir:
        output_path = os.path.join(tmpdir, "output.docx")
        write_docx(results, output_path)

        assert os.path.exists(output_path)

        # Verify the DOCX can be reopened
        doc = docx.Document(output_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        assert len(paragraphs) == 2
        assert "नमस्ते दुनिया" in paragraphs[0]
        assert "यह एक परीक्षण है" in paragraphs[1]

def test_write_docx_unicode_tamil():
    results = [
        TranslationResult("chunk-1", "வணக்கம் உலகம்", "tamil"),
    ]

    with tempfile.TemporaryDirectory() as tmpdir:
        output_path = os.path.join(tmpdir, "output.docx")
        write_docx(results, output_path)

        doc = docx.Document(output_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        assert "வணக்கம்" in paragraphs[0]
