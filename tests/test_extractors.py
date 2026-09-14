import os
import tempfile
import pytest
import docx
import fitz # PyMuPDF
from app.extractors import extract_document
from app.exceptions import CIPAMError

def create_temp_pdf(content_pages):
    fd, path = tempfile.mkstemp(suffix=".pdf")
    os.close(fd)
    
    doc = fitz.open()
    for content in content_pages:
        page = doc.new_page()
        page.insert_text(fitz.Point(50, 50), content)
    doc.save(path)
    doc.close()
    return path

def create_temp_docx(paragraphs, headings):
    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    
    doc = docx.Document()
    for i, p in enumerate(paragraphs):
        if i in headings:
            doc.add_heading(p, level=1)
        else:
            doc.add_paragraph(p)
    doc.save(path)
    return path

def test_extract_pdf():
    pdf_path = create_temp_pdf(["Page 1 content", "Page 2 content"])
    try:
        doc = extract_document(pdf_path)
        assert doc.source_format == "pdf"
        assert len(doc.text_blocks) == 2
        assert "Page 1" in doc.text_blocks[0].text
        assert doc.text_blocks[0].page_number == 1
        assert "Page 2" in doc.text_blocks[1].text
        assert doc.text_blocks[1].page_number == 2
    finally:
        os.remove(pdf_path)

def test_extract_docx():
    docx_path = create_temp_docx(["Main Title", "Paragraph 1", "Paragraph 2"], headings=[0])
    try:
        doc = extract_document(docx_path)
        assert doc.source_format == "docx"
        assert len(doc.text_blocks) == 3
        
        assert doc.text_blocks[0].text == "Main Title"
        assert doc.text_blocks[0].block_type == "heading"
        
        assert doc.text_blocks[1].text == "Paragraph 1"
        assert doc.text_blocks[1].block_type == "paragraph"
    finally:
        os.remove(docx_path)

def test_extract_unsupported_type():
    with pytest.raises(CIPAMError, match="Unsupported file type"):
        extract_document("test.csv")

def test_extract_missing_pdf():
    with pytest.raises(CIPAMError, match="File not found"):
        extract_document("missing.pdf")
