import os
import docx
from app.models import Document, TextBlock
from app.exceptions import CIPAMError, ExtractionError

def extract_docx(file_path: str) -> Document:
    if not os.path.exists(file_path):
        raise CIPAMError(f"File not found: {file_path}")

    filename = os.path.basename(file_path)
    doc = Document(source_filename=filename, source_format="docx")

    try:
        doc_obj = docx.Document(file_path)
    except Exception as e:
        raise ExtractionError(f"Failed to open DOCX file {file_path}: {e}")

    index = 0
    for para in doc_obj.paragraphs:
        text = para.text.strip()
        if not text:
            continue
            
        # Try to detect if it's a heading
        block_type = "paragraph"
        if para.style and para.style.name.startswith("Heading"):
            block_type = "heading"
            
        block = TextBlock(
            text=text,
            block_type=block_type,
            index=index
        )
        doc.text_blocks.append(block)
        index += 1

    return doc
